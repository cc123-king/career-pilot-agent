from collections.abc import Iterator
from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.document import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from career_pilot.multimodal.models import (
    ExtractedDocument,
    ExtractedPage,
)
from career_pilot.multimodal.routing import route_document


def _iter_document_blocks(
    document: WordDocument,
) -> Iterator[Paragraph | Table]:
    """Yield paragraphs and tables in their original order."""

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _extract_table_rows(table: Table) -> list[str]:
    """Convert Word table rows into readable text."""

    extracted_rows: list[str] = []

    for row in table.rows:
        cell_texts = [
            " ".join(
                paragraph.text.strip()
                for paragraph in cell.paragraphs
                if paragraph.text.strip()
            )
            for cell in row.cells
        ]

        row_text = " | ".join(
            text for text in cell_texts if text
        )

        if row_text:
            extracted_rows.append(row_text)

    return extracted_rows


def load_word_document(
    path: str | Path,
) -> ExtractedDocument:
    """Extract paragraphs and tables from a DOCX file."""

    document_path, source_type = route_document(path)

    if source_type != "docx":
        raise ValueError("Word 加载器只支持 .docx 文件")

    try:
        word_document = Document(document_path)
    except (PackageNotFoundError, BadZipFile) as exc:
        raise ValueError(
            "Word 文件损坏或格式无效"
        ) from exc

    content_blocks: list[str] = []

    for block in _iter_document_blocks(word_document):
        if isinstance(block, Paragraph):
            paragraph_text = block.text.strip()

            if paragraph_text:
                content_blocks.append(paragraph_text)
        else:
            content_blocks.extend(
                _extract_table_rows(block)
            )

    if not content_blocks:
        raise ValueError("Word 文件没有可提取的文字")

    content = "\n\n".join(content_blocks)

    return ExtractedDocument(
        filename=document_path.name,
        source_path=str(document_path.resolve()),
        source_type="docx",
        pages=[
            ExtractedPage(
                page_number=1,
                text=content,
                extraction_method="native_text",
            )
        ],
    )