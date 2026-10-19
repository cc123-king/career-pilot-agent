from pathlib import Path

import pytest
from docx import Document

from career_pilot.multimodal.word_loader import (
    load_word_document,
)


def test_word_loader_extracts_paragraphs_and_tables(
    tmp_path: Path,
) -> None:
    document_path = tmp_path / "resume.docx"

    word_document = Document()
    word_document.add_heading("示例候选人个人简历", level=1)

    table = word_document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "示例候选人"
    table.cell(1, 0).text = "技能"
    table.cell(1, 1).text = "Python、RAG、FastAPI"

    word_document.add_paragraph(
        "项目经历：CareerPilot Agent"
    )
    word_document.save(document_path)

    document = load_word_document(document_path)

    assert document.filename == "resume.docx"
    assert document.source_type == "docx"
    assert document.source_path == str(
        document_path.resolve()
    )
    assert document.pages[0].extraction_method == "native_text"
    assert "示例候选人个人简历" in document.content
    assert "姓名 | 示例候选人" in document.content
    assert "技能 | Python、RAG、FastAPI" in document.content
    assert "CareerPilot Agent" in document.content

    assert document.content.index(
        "示例候选人个人简历"
    ) < document.content.index(
        "姓名 | 示例候选人"
    )
    assert document.content.index(
        "姓名 | 示例候选人"
    ) < document.content.index(
        "CareerPilot Agent"
    )


def test_word_loader_rejects_non_docx_file(
    tmp_path: Path,
) -> None:
    document_path = tmp_path / "resume.txt"
    document_path.write_text(
        "普通文本",
        encoding="utf-8",
    )

    with pytest.raises(
        ValueError,
        match="只支持 .docx",
    ):
        load_word_document(document_path)


def test_word_loader_rejects_corrupt_docx(
    tmp_path: Path,
) -> None:
    document_path = tmp_path / "broken.docx"
    document_path.write_bytes(b"not a real Word file")

    with pytest.raises(
        ValueError,
        match="损坏或格式无效",
    ):
        load_word_document(document_path)


def test_word_loader_rejects_empty_document(
    tmp_path: Path,
) -> None:
    document_path = tmp_path / "empty.docx"

    word_document = Document()
    word_document.save(document_path)

    with pytest.raises(
        ValueError,
        match="没有可提取的文字",
    ):
        load_word_document(document_path)